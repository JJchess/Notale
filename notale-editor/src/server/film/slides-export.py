"""Package captured film frames. Slide contents are raster images, with source times in notes."""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / '.local/motale-film/python'))
manifest_path, destination = map(Path, sys.argv[1:3])
manifest = json.loads(manifest_path.read_text())
width, height = manifest['width'], manifest['height']
images = [manifest_path.parent / item['file'] for item in manifest['slides']]
if destination.suffix == '.pptx':
    from pptx import Presentation
    presentation = Presentation()
    factor = 9144000 / max(width, height)
    slide_width, slide_height = round(width * factor), round(height * factor)
    if min(slide_width, slide_height) < 914400:
        raise ValueError('PPTX 不支持超过 10:1 的画幅，请用 PDF 或 PNG 集')
    presentation.slide_width = slide_width
    presentation.slide_height = slide_height
    presentation.core_properties.title = manifest['title']
    presentation.core_properties.subject = 'Notale film scene stills'
    for image, item in zip(images, manifest['slides']):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_picture(str(image), 0, 0, width=slide_width, height=slide_height)
        shape.name = item['title']
        slide.notes_slide.notes_text_frame.text = f"{item['title']}\nSource frame: {item['frame']}\nTime: {item['seconds']:.6f} seconds\nRaster snapshot of the film compositor."
    presentation.save(str(destination))
    check = Presentation(str(destination))
    assert len(check.slides) == len(images)
else:
    raise ValueError('Unsupported slide output')
print(json.dumps({'slides': len(images), 'width': width, 'height': height, 'raster': True}))
